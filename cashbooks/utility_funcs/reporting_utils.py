from datetime import timedelta
from django.db.models import Sum, Count
from django.utils import timezone
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.enums import TA_CENTER
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from io import BytesIO


def get_date_range(period_type, custom_start=None, custom_end=None):
    """
    Get start and end dates based on period type.
    period_type: 'daily', 'weekly', 'monthly', 'custom'
    """
    today = timezone.now().date()
    
    if period_type == 'daily':
        return today, today
    elif period_type == 'weekly':
        start_date = today - timedelta(days=today.weekday())  # Monday
        end_date = start_date + timedelta(days=6)  # Sunday
        return start_date, end_date
    elif period_type == 'monthly':
        start_date = today.replace(day=1)
        # Get last day of month
        if today.month == 12:
            end_date = today.replace(day=31)
        else:
            end_date = (today.replace(month=today.month + 1, day=1) - timedelta(days=1))
        return start_date, end_date
    elif period_type == 'custom':
        if not custom_start or not custom_end:
            raise ValueError("Custom date range requires start_date and end_date")
        return custom_start, custom_end
    else:
        raise ValueError("Invalid period_type")


def generate_pdf_report(cashbook, entries, start_date, end_date):
    """
    Generate PDF report for cashbook entries.
    Returns BytesIO buffer with PDF content.
    """
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.5*inch, bottomMargin=0.5*inch)
    elements = []
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'CustomTitle',
        parent=styles['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#1a1a1a'),
        spaceAfter=12,
        alignment=TA_CENTER
    )
    
    heading_style = ParagraphStyle(
        'CustomHeading',
        parent=styles['Heading2'],
        fontSize=14,
        textColor=colors.HexColor('#333333'),
        spaceAfter=10,
        spaceBefore=10
    )
    
    # Title
    title = Paragraph(f"<b>Financial Report: {cashbook.book_name}</b>", title_style)
    elements.append(title)
    elements.append(Spacer(1, 0.2*inch))
    
    # Report Info
    info_data = [
        ['Report Period:', f"{start_date.strftime('%d %b %Y')} to {end_date.strftime('%d %b %Y')}"],
        ['Generated On:', timezone.now().strftime('%d %b %Y, %I:%M %p')],
        ['Owner:', cashbook.owner.email],
    ]
    info_table = Table(info_data, colWidths=[2*inch, 4*inch])
    info_table.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 10),
        ('TEXTCOLOR', (0, 0), (0, -1), colors.HexColor('#555555')),
        ('ALIGN', (0, 0), (0, -1), 'LEFT'),
        ('ALIGN', (1, 0), (1, -1), 'LEFT'),
    ]))
    elements.append(info_table)
    elements.append(Spacer(1, 0.3*inch))
    
    # Summary Section
    cash_in_total = entries.filter(entry_type='cash_in').aggregate(total=Sum('amount'))['total'] or 0
    cash_out_total = entries.filter(entry_type='cash_out').aggregate(total=Sum('amount'))['total'] or 0
    net_balance = cash_in_total - cash_out_total
    
    elements.append(Paragraph("<b>Summary</b>", heading_style))
    summary_data = [
        ['Total Cash In', f"৳ {cash_in_total:,.2f}"],
        ['Total Cash Out', f"৳ {cash_out_total:,.2f}"],
        ['Net Balance', f"৳ {net_balance:,.2f}"],
    ]
    summary_table = Table(summary_data, colWidths=[3*inch, 3*inch])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e8f4f8')),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#fff4e6')),
        ('BACKGROUND', (0, 2), (-1, 2), colors.HexColor('#e8f8e8')),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 11),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 0.3*inch))
    
    # Entries Section
    elements.append(Paragraph("<b>Transaction Details</b>", heading_style))
    
    if entries.exists():
        # Table headers
        entry_data = [['Date', 'Type', 'Title', 'Category', 'Payment', 'Amount']]
        
        # Table rows
        for entry in entries:
            entry_data.append([
                entry.entry_date.strftime('%d %b %Y'),
                'Cash In' if entry.entry_type == 'cash_in' else 'Cash Out',
                entry.title or '-',
                entry.category.category_name if entry.category else '-',
                entry.payment_method.payment_method_name if entry.payment_method else '-',
                f"৳ {entry.amount:,.2f}"
            ])
        
        entry_table = Table(entry_data, colWidths=[1*inch, 1*inch, 1.5*inch, 1.2*inch, 1.2*inch, 1*inch])
        entry_table.setStyle(TableStyle([
            # Header styling
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4a90e2')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            
            # Body styling
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('ALIGN', (0, 1), (0, -1), 'CENTER'),  # Date
            ('ALIGN', (1, 1), (1, -1), 'CENTER'),  # Type
            ('ALIGN', (-1, 1), (-1, -1), 'RIGHT'),  # Amount
            
            # Grid
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f9f9f9')]),
        ]))
        elements.append(entry_table)
    else:
        elements.append(Paragraph("<i>No transactions found for this period.</i>", styles['Normal']))
    
    # Build PDF
    doc.build(elements)
    buffer.seek(0)
    return buffer


def generate_excel_report(cashbook, entries, start_date, end_date):
    """
    Generate Excel report for cashbook entries.
    Returns BytesIO buffer with Excel content.
    """
    buffer = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Financial Report"
    
    # Define styles
    header_fill = PatternFill(start_color="4A90E2", end_color="4A90E2", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    title_font = Font(bold=True, size=16)
    bold_font = Font(bold=True, size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Title
    ws.merge_cells('A1:F1')
    ws['A1'] = f"Financial Report: {cashbook.book_name}"
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    
    # Report Info
    ws['A3'] = "Report Period:"
    ws['A3'].font = bold_font
    ws['B3'] = f"{start_date.strftime('%d %b %Y')} to {end_date.strftime('%d %b %Y')}"
    
    ws['A4'] = "Generated On:"
    ws['A4'].font = bold_font
    ws['B4'] = timezone.now().strftime('%d %b %Y, %I:%M %p')
    
    ws['A5'] = "Owner:"
    ws['A5'].font = bold_font
    ws['B5'] = cashbook.owner.email
    
    # Summary Section
    cash_in_total = entries.filter(entry_type='cash_in').aggregate(total=Sum('amount'))['total'] or 0
    cash_out_total = entries.filter(entry_type='cash_out').aggregate(total=Sum('amount'))['total'] or 0
    net_balance = cash_in_total - cash_out_total
    
    ws['A7'] = "Summary"
    ws['A7'].font = Font(bold=True, size=14)
    
    ws['A8'] = "Total Cash In"
    ws['A8'].font = bold_font
    ws['B8'] = float(cash_in_total)
    ws['B8'].number_format = '৳#,##0.00'
    
    ws['A9'] = "Total Cash Out"
    ws['A9'].font = bold_font
    ws['B9'] = float(cash_out_total)
    ws['B9'].number_format = '৳#,##0.00'
    
    ws['A10'] = "Net Balance"
    ws['A10'].font = bold_font
    ws['B10'] = float(net_balance)
    ws['B10'].number_format = '৳#,##0.00'
    
    # Entries Section
    ws['A12'] = "Transaction Details"
    ws['A12'].font = Font(bold=True, size=14)
    
    # Headers
    headers = ['Date', 'Type', 'Title', 'Category', 'Payment Method', 'Amount']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=14, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    
    # Data rows
    row_num = 15
    for entry in entries:
        ws.cell(row=row_num, column=1, value=entry.entry_date.strftime('%d %b %Y')).border = border
        ws.cell(row=row_num, column=2, value='Cash In' if entry.entry_type == 'cash_in' else 'Cash Out').border = border
        ws.cell(row=row_num, column=3, value=entry.title or '-').border = border
        ws.cell(row=row_num, column=4, value=entry.category.category_name if entry.category else '-').border = border
        ws.cell(row=row_num, column=5, value=entry.payment_method.payment_method_name if entry.payment_method else '-').border = border
        
        amount_cell = ws.cell(row=row_num, column=6, value=float(entry.amount))
        amount_cell.number_format = '৳#,##0.00'
        amount_cell.border = border
        
        row_num += 1
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 20
    ws.column_dimensions['F'].width = 15
    
    # Save to buffer
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def get_analytics_data(cashbook, start_date=None, end_date=None):
    """
    Get analytics data for cashbook entries.
    Returns data suitable for Flutter charts (pie charts, bar charts, etc.)
    """
    # Filter entries by date range if provided
    entries = cashbook.entry_set.all()
    if start_date and end_date:
        entries = entries.filter(entry_date__range=[start_date, end_date])
    
    # Category-wise spending (Cash Out only)
    category_data = entries.filter(
        entry_type='cash_out',
        category__isnull=False,
        category__is_deleted=False
    ).values(
        'category__category_name'
    ).annotate(
        total_amount=Sum('amount'),
        count=Count('id')
    ).order_by('-total_amount')
    
    category_analytics = {
        'labels': [item['category__category_name'] for item in category_data],
        'values': [float(item['total_amount']) for item in category_data],
        'counts': [item['count'] for item in category_data],
    }
    
    # Category-wise income (Cash In only)
    category_income_data = entries.filter(
        entry_type='cash_in',
        category__isnull=False,
        category__is_deleted=False
    ).values(
        'category__category_name'
    ).annotate(
        total_amount=Sum('amount'),
        count=Count('id')
    ).order_by('-total_amount')
    
    category_income_analytics = {
        'labels': [item['category__category_name'] for item in category_income_data],
        'values': [float(item['total_amount']) for item in category_income_data],
        'counts': [item['count'] for item in category_income_data],
    }
    
    # Payment method-wise analysis (both cash in and out)
    payment_method_data = entries.filter(
        payment_method__isnull=False,
        payment_method__is_deleted=False
    ).values(
        'payment_method__payment_method_name',
        'entry_type'
    ).annotate(
        total_amount=Sum('amount'),
        count=Count('id')
    ).order_by('payment_method__payment_method_name', 'entry_type')
    
    # Organize payment method data by type
    payment_analytics = {}
    for item in payment_method_data:
        method = item['payment_method__payment_method_name']
        entry_type = item['entry_type']
        
        if method not in payment_analytics:
            payment_analytics[method] = {
                'cash_in': 0,
                'cash_out': 0,
                'cash_in_count': 0,
                'cash_out_count': 0,
            }
        
        if entry_type == 'cash_in':
            payment_analytics[method]['cash_in'] = float(item['total_amount'])
            payment_analytics[method]['cash_in_count'] = item['count']
        else:
            payment_analytics[method]['cash_out'] = float(item['total_amount'])
            payment_analytics[method]['cash_out_count'] = item['count']
    
    # Daily trend (last 30 days or specified range)
    if not start_date or not end_date:
        end_date = timezone.now().date()
        start_date = end_date - timedelta(days=30)
    
    daily_data = entries.filter(
        entry_date__range=[start_date, end_date]
    ).values('entry_date', 'entry_type').annotate(
        total_amount=Sum('amount')
    ).order_by('entry_date')
    
    # Organize daily data
    daily_analytics = {}
    for item in daily_data:
        date_str = item['entry_date'].strftime('%Y-%m-%d')
        if date_str not in daily_analytics:
            daily_analytics[date_str] = {'cash_in': 0, 'cash_out': 0}
        
        if item['entry_type'] == 'cash_in':
            daily_analytics[date_str]['cash_in'] = float(item['total_amount'])
        else:
            daily_analytics[date_str]['cash_out'] = float(item['total_amount'])
    
    # Calculate totals
    cash_in_total = entries.filter(entry_type='cash_in').aggregate(total=Sum('amount'))['total'] or 0
    cash_out_total = entries.filter(entry_type='cash_out').aggregate(total=Sum('amount'))['total'] or 0
    
    return {
        'summary': {
            'total_cash_in': float(cash_in_total),
            'total_cash_out': float(cash_out_total),
            'net_balance': float(cash_in_total - cash_out_total),
            'total_transactions': entries.count(),
        },
        'category_spending': category_analytics,
        'category_income': category_income_analytics,
        'payment_methods': payment_analytics,
        'daily_trend': daily_analytics,
        'date_range': {
            'start': start_date.strftime('%Y-%m-%d'),
            'end': end_date.strftime('%Y-%m-%d'),
        }
    }